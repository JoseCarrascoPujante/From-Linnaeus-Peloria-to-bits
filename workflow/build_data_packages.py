#!/usr/bin/env python3
"""Supplementary Data S1, S2 and the Peloria entries of Data S7 from the reconstructed matrices.

The default mode checks that the deposited packages contain exactly what the
reconstructed Lcyc/LvHIRZ count matrices imply and exits non-zero otherwise;
--write rebuilds them (04_Supplementary_Data and the Peloria input of
06_EpiBit_Benchmark).
Zip members keep their order and a fixed timestamp so rebuilt packages are
byte-reproducible.
"""
from __future__ import annotations
import argparse, csv, io, math, re, sys, zipfile
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'02_FASTA_to_Counts_Reconstruction'/'code'))
from config import ALPHA, MIN_IDENTITY, MIN_OVERLAP, SEED, PERMUTATIONS  # noqa: E402

PRE = ROOT/'02_FASTA_to_Counts_Reconstruction'/'outputs'/'peloria_reconstructed'
SUPP = ROOT/'04_Supplementary_Data'
BENCH = ROOT/'06_EpiBit_Benchmark'/'inputs'/'Peloria_Lcyc.csv'
ZIP_TIME = (2026, 10, 27, 0, 0, 0)
CONTEXTS = ['CG', 'CHG', 'CHH']
ALPHAS = [0.1, 0.25, 0.5, 1.0, 2.0]
THRESHOLDS = [0.5, 0.8, 0.9, 0.95]
LOCI = {'Lcyc': ('Supplementary_Data_S1_Lcyc', 'Data_S1_Lcyc.zip', 'S7', 'S1'),
        'LvHIRZ': ('Supplementary_Data_S2_LvHIRZ', 'Data_S2_LvHIRZ.zip', 'S8', 'S2')}

# ---------------------------------------------------------------- site metrics
def site_table(matrix: pd.DataFrame, alpha: float = ALPHA) -> pd.DataFrame:
    d = matrix.copy()
    p = (d['WT methylated'] + alpha)/(d['WT total'] + 2*alpha)
    q = (d['Peloria methylated'] + alpha)/(d['Peloria total'] + 2*alpha)
    h = lambda x: -(x*np.log2(x) + (1 - x)*np.log2(1 - x))
    kl = lambda a, b: a*np.log2(a/b) + (1 - a)*np.log2((1 - a)/(1 - b))
    m = (p + q)/2
    d['p'], d['q'] = p, q
    d['H_WT'], d['H_Peloria'] = h(p), h(q)
    d['KL'], d['KL_rev'] = kl(q, p), kl(p, q)
    d['D_sym'] = (d.KL + d.KL_rev)/2
    d['JSD'] = (kl(p, m) + kl(q, m))/2
    return d

def thresholds(dsym: np.ndarray) -> list[int]:
    # Smallest number of top-ranked sites whose cumulative D_sym reaches each fraction.
    cum = np.cumsum(np.sort(dsym)[::-1])/dsym.sum()
    return [int(np.argmax(cum >= t)) + 1 for t in THRESHOLDS]

def with_site_ids(matrix: pd.DataFrame) -> pd.DataFrame:
    d = matrix.copy()
    d['site_id'] = [f'{f}_C{i:03d}' for i, f in enumerate(d.fragment, 1)]
    return d

# ------------------------------------------------------------- S1/S2 contents
def matrix_csv(d: pd.DataFrame) -> bytes:
    out = d[['site_id', 'fragment', 'position', 'context', 'WT methylated', 'WT total',
             'Peloria methylated', 'Peloria total']].copy()
    out['WT p'] = d.p.map('{:.3f}'.format); out['Peloria q'] = d.q.map('{:.3f}'.format)
    return out.to_csv(index=False, lineterminator='\n').encode()

def context_summary_csv(d: pd.DataFrame) -> bytes:
    groups = [(c, d[d.context == c]) for c in CONTEXTS + ['C_unknown'] if (d.context == c).any()]
    groups.append(('Total', d))
    rows = []
    for name, g in groups:
        rows.append([name, len(g), g.p.mean(), g.q.mean(), g.H_WT.sum(), g.H_Peloria.sum(), g.KL.sum(),
                     g.KL_rev.sum(), g.D_sym.sum(), g.JSD.sum(), 100*g.KL.sum()/d.KL.sum(),
                     100*g.JSD.sum()/d.JSD.sum()])
    buf = io.StringIO(); w = csv.writer(buf, lineterminator='\n')
    w.writerow(['context', 'N', 'WT_mean', 'Peloria_mean', 'H_WT', 'H_Peloria', 'KL_Peloria_WT',
                'KL_WT_Peloria', 'D_sym', 'JSD', 'KL_%', 'JSD_%'])
    for r in rows:
        w.writerow([r[0], r[1]] + [f'{x:.3f}' for x in r[2:]])
    return buf.getvalue().encode()

def offsets_by_count(text: str) -> str:
    # Workbook convention: offsets listed by decreasing read count.
    pairs = [tuple(map(int, x.split(':'))) for x in text.split('; ')]
    return '; '.join(f'{o}:{n}' for o, n in sorted(pairs, key=lambda x: (-x[1], x[0])))

def modal_offset(text: str) -> int:
    pairs = [tuple(map(int, x.split(':'))) for x in text.split('; ')]
    return max(pairs, key=lambda x: (x[1], -abs(x[0])))[0]

ALIGNMENT_RULE = (f'Each read was aligned without gaps at every offset with at least {MIN_OVERLAP} '
                  f'overlapping bases using bisulfite-compatible C/T matching, and accepted if its best '
                  f'identity was at least {MIN_IDENTITY}.')

def fill_rows(ws, rows):
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v.item() if hasattr(v, 'item') else v)

def lcyc_workbook(template: bytes, d: pd.DataFrame, qc: pd.DataFrame) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(template))
    ws = wb['README']
    readme = [
        ('Lcyc cytosine-level methylation matrix from bisulfite/NGS reads', None),
        ('Source', 'Targeted bisulfite/NGS FASTA reads provided by Prof. Ann E. Ehrenhofer-Murray (Ehrenhofer-Murray, 2021).'),
        ('Samples', 'File suffix _8 = wild-type Linaria vulgaris; file suffix _9 = Peloria mutant'),
        ('Lcyc fragments analyzed', 'LCYC_F3_R3, LCYC_peloria_fragment3, LCYC_peloria_fragment4'),
        ('Core output', 'Cytosine_matrix contains methylated and total read counts per assayed cytosine.'),
        ('Counting rule', 'At reference C positions, read C was counted as methylated and read T as unmethylated; other calls were excluded and are reported in Alignment_QC. LCYC_peloria_fragment4 reads were reverse-complemented before counting because they are in reverse-complement orientation relative to the reference fragment.'),
        ('Alignment filter', ALIGNMENT_RULE),
        ('Pseudocount in Derived_metrics', 'alpha = 0.5, so p = (m + alpha)/(n + 2 alpha).'),
    ]
    ws.delete_rows(1, ws.max_row)
    for row in readme:
        ws.append(list(row))
    base = [[r.site_id, r.fragment, int(r.position), r.context, int(r['WT methylated']), int(r['WT total']),
             int(r['Peloria methylated']), int(r['Peloria total'])] for _, r in d.iterrows()]
    fill_rows(wb['Cytosine_matrix'], base)
    fill_rows(wb['Derived_metrics'], [b + [r.p, r.q, r['WT methylated']/r['WT total'],
                                           r['Peloria methylated']/r['Peloria total'], r.H_WT, r.H_Peloria,
                                           r.KL, r.KL_rev, r.D_sym, r.JSD]
                                      for b, (_, r) in zip(base, d.iterrows())])
    summ = []
    for kind, key, groups in [('context', 'context', CONTEXTS), ('fragment', 'fragment', list(dict.fromkeys(d.fragment))),
                              ('overall', None, ['ALL'])]:
        for g in groups:
            s = d if key is None else d[d[key] == g]
            summ.append([kind, g, len(s), int(s['WT methylated'].sum()), int(s['WT total'].sum()),
                         int(s['Peloria methylated'].sum()), int(s['Peloria total'].sum()), s.p.mean(), s.q.mean(),
                         s.H_WT.sum(), s.H_Peloria.sum(), s.KL.sum(), s.JSD.sum()])
    fill_rows(wb['Summary'], summ)
    ws = wb['Alignment_QC']
    ws.cell(1, 11, 'modal_offset')
    tab = ws.tables['AlignmentQCTable']; tab.tableColumns[10].name = 'modal_offset'
    rows = []
    for _, r in qc[qc.locus == 'Lcyc'].iterrows():
        rows.append([r.fragment, r.condition, r.file, int(r.raw_reads), int(r.accepted_reads), int(r.rejected_reads),
                     r.acceptance_rate, r.mean_identity, r.min_identity, r.max_identity,
                     modal_offset(r.observed_offsets), offsets_by_count(r.observed_offsets),
                     r.orientations.startswith('reverse'), int(r.ignored_non_CT_at_C_sites)])
    fill_rows(ws, rows)
    rk = d.sort_values('D_sym', ascending=False, kind='stable').reset_index(drop=True)
    frac = rk.D_sym/rk.D_sym.sum()
    fill_rows(wb['Dimensionality'], [[i + 1, r.site_id, r.fragment, int(r.position), r.context, r.D_sym, frac[i],
                                      frac[:i + 1].sum()] for i, r in rk.iterrows()])
    return save_workbook(wb)

def lvhirz_workbook(template: bytes, qc: pd.DataFrame) -> bytes:
    # Matrix sheets are unchanged; QC now describes accepted reads only.
    wb = openpyxl.load_workbook(io.BytesIO(template))
    ws = wb['README']
    for row in ws.iter_rows(min_row=1):
        if row[0].value == 'Reference fragment':
            row[1].value = 'LvHIRZ_2 reference fragment (reference_sequences.fasta in the study repository).'
        if row[0].value == 'Alignment rule':
            row[1].value = ALIGNMENT_RULE.replace('without gaps', 'without gaps, in reverse-complement orientation,')
    ws = wb['Alignment_QC']
    for i, (_, r) in enumerate(qc[qc.locus == 'LvHIRZ'].iterrows(), 2):
        ori = '; '.join(('rc' if k == 'reverse' else 'direct') + ':' + v
                        for k, v in (x.split(':') for x in r.orientations.split('; ')))
        vals = [r.fragment, r.condition, r.file, r.raw_reads, r.accepted_reads, r.rejected_reads, r.acceptance_rate,
                r.mean_identity, r.min_identity, r.max_identity, ori, offsets_by_count(r.observed_offsets),
                r.ignored_non_CT_at_C_sites]
        for c, v in enumerate(vals, 1):
            ws.cell(i, c, str(v))  # this workbook stores every value as text
    return save_workbook(wb)

def save_workbook(wb) -> bytes:
    import datetime
    t = datetime.datetime(*ZIP_TIME)
    wb.properties.created = t; wb.properties.modified = t
    buf = io.BytesIO(); wb.save(buf)
    # Normalise the xlsx container itself so identical content gives identical bytes.
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue())); out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for it in src.infolist():
            data = src.read(it.filename)
            if it.filename == 'docProps/core.xml':
                # openpyxl stamps the save time into 'modified'; pin it like 'created'.
                data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*', rb'\g<1>' + t.strftime('%Y-%m-%dT%H:%M:%SZ').encode(), data)
            z.writestr(zipfile.ZipInfo(it.filename, ZIP_TIME), data, zipfile.ZIP_DEFLATED)
    return out.getvalue()

# --------------------------------------------------------- Data S7 (Peloria)
def architecture_row(d: pd.DataFrame) -> dict:
    n = len(d); tot = d.KL.sum()
    p = {c: (d.context == c).sum()/n for c in CONTEXTS}
    q = {c: d.KL[d.context == c].sum()/tot for c in CONTEXTS}
    e = {c: q[c]/p[c] for c in CONTEXTS}
    nt = thresholds(d.D_sym.to_numpy())
    row = {'system': 'Peloria', 'N': n, 'KL_bits': tot, 'Dsym_bits': d.D_sym.sum(), 'JSD_bits': d.JSD.sum()}
    row.update({f'p_{c}': p[c] for c in CONTEXTS}); row.update({f'q_{c}': q[c] for c in CONTEXTS})
    row.update({f'E_{c}': e[c] for c in CONTEXTS})
    row.update({f'N{int(t*100)}': k for t, k in zip(THRESHOLDS, nt)})
    row.update({f'R{int(t*100)}': k/n for t, k in zip(THRESHOLDS, nt)})
    row['X_nonCG_dominance'] = math.log2((q['CHG'] + q['CHH'])/q['CG'])
    row['Y_log2_ECHG'] = math.log2(e['CHG'])
    row['Z_log2_ECG'] = math.log2(e['CG'])
    row['context_breadth_B'] = -sum(q[c]*math.log2(q[c]) for c in CONTEXTS)/math.log2(3)
    row['ILR1_CG_vs_nonCG'] = math.sqrt(2/3)*math.log(q['CG']/math.sqrt(q['CHG']*q['CHH']))
    row['ILR2_CHG_vs_CHH'] = math.sqrt(1/2)*math.log(q['CHG']/q['CHH'])
    return row

def sensitivity_rows(matrix: pd.DataFrame) -> list[dict]:
    rows = []
    for a in ALPHAS:
        d = site_table(matrix, a); tot = d.KL.sum()
        nt = thresholds(d.D_sym.to_numpy())
        rows.append({'system': 'Peloria', 'alpha': a, 'KL_bits': tot, 'Dsym_bits': d.D_sym.sum(), 'JSD_bits': d.JSD.sum(),
                     'q_CHG': d.KL[d.context == 'CHG'].sum()/tot, 'q_nonCG': d.KL[d.context != 'CG'].sum()/tot,
                     **{f'N{int(t*100)}': k for t, k in zip(THRESHOLDS, nt)}})
    return rows

def permutation_rows(d: pd.DataFrame) -> list[dict]:
    """Context-label permutation: labels reshuffled across sites, site-wise KL and context sizes fixed."""
    kl = d.KL.to_numpy(); tot = kl.sum()
    codes = np.array([CONTEXTS.index(c) for c in d.context]); n = len(codes)
    rng = np.random.default_rng(SEED)
    null = np.empty((PERMUTATIONS, 3))
    for i in range(PERMUTATIONS):
        perm = rng.permutation(codes)
        null[i] = np.bincount(perm, weights=kl, minlength=3)/tot
    rows = []
    for j, c in enumerate(CONTEXTS):
        obs = kl[codes == j].sum()/tot; frac = (codes == j).mean(); x = null[:, j]
        rows.append({'system': 'Peloria', 'context': c, 'N_context': int((codes == j).sum()), 'N_total': n,
                     'site_fraction': frac, 'observed_divergence_fraction': obs, 'enrichment': obs/frac,
                     'null_mean': x.mean(), 'null_p0_1': np.quantile(x, 0.001), 'null_p99_9': np.quantile(x, 0.999),
                     # (b + 1)/(N_perm + 1), as defined in the manuscript methods.
                     'p_upper': ((x >= obs).sum() + 1)/(PERMUTATIONS + 1),
                     'p_lower': ((x <= obs).sum() + 1)/(PERMUTATIONS + 1),
                     'permutations': PERMUTATIONS, 'seed': SEED})
    return rows

def replace_peloria_rows(old_csv: bytes, rows: pd.DataFrame) -> bytes:
    # Rows of the other systems are kept verbatim; only the Peloria rows are regenerated.
    eol = '\r\n' if b'\r\n' in old_csv else '\n'  # keep each file's original line endings
    lines = old_csv.decode().splitlines(); cols = lines[0].split(',')
    i = cols.index('system') if 'system' in cols else cols.index('system1')
    kept = [l for l in lines[1:] if l.split(',')[i] != 'Peloria']
    fresh = rows[cols].to_csv(index=False, header=False, lineterminator='\n').splitlines()
    return (eol.join([lines[0]] + fresh + kept) + eol).encode()

def s7_members(old: dict[str, bytes], d: pd.DataFrame, matrix: pd.DataFrame) -> dict[str, bytes]:
    pre = 'Data_S7_architecture_space/'
    rd = lambda name: pd.read_csv(io.BytesIO(old[pre + name]), float_precision='round_trip')
    arch_row = pd.DataFrame([architecture_row(d)])
    xy = pd.concat([arch_row, rd('architecture_source_and_coordinates.csv').query("system != 'Peloria'")]).set_index('system')
    dist = rd('pairwise_primary_space_distances.csv').query("system1 == 'Peloria'").copy()
    dist['distance_XY'] = [math.hypot(xy.loc[a].X_nonCG_dominance - xy.loc[b].X_nonCG_dominance,
                                      xy.loc[a].Y_log2_ECHG - xy.loc[b].Y_log2_ECHG) for a, b in zip(dist.system1, dist.system2)]
    perm = pd.DataFrame(permutation_rows(d))
    new = dict(old)
    for name, rows in [('architecture_source_and_coordinates.csv', arch_row),
                       ('pairwise_primary_space_distances.csv', dist),
                       ('pseudocount_sensitivity_all_systems.csv', pd.DataFrame(sensitivity_rows(matrix))),
                       ('context_label_permutation_summary.csv', perm)]:
        new[pre + name] = replace_peloria_rows(old[pre + name], rows)
    return new

# ------------------------------------------------------------------ zip I/O
def read_zip(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}

def write_zip(path: Path, members: dict[str, bytes]) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(zipfile.ZipInfo(name, ZIP_TIME), data, zipfile.ZIP_DEFLATED)
    path.write_bytes(buf.getvalue())

def sheet_values(xlsx: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}

def same(a: bytes, b: bytes, name: str) -> bool:
    if name.endswith('.xlsx'):
        return sheet_values(a) == sheet_values(b)
    return a == b

# --------------------------------------------------------------------- main
def build() -> dict[Path, dict[str, bytes]]:
    qc = pd.read_csv(PRE/'alignment_QC_reconstructed.csv')
    packages = {}
    tables = {}
    for locus, (folder, zname, mtab, ctab) in LOCI.items():
        matrix = pd.read_csv(PRE/f'{locus}_cytosine_matrix_reconstructed.csv')
        d = with_site_ids(site_table(matrix)); tables[locus] = (matrix, d)
        old = read_zip(SUPP/zname); pre = folder + '/'
        xname = pre + f'{locus}_empirical_cytosine_methylation_matrix.xlsx'
        new = {pre + f'Supplementary_Table_{ctab}_{locus}_context_summary.csv': context_summary_csv(d),
               pre + f'Supplementary_Table_{mtab}_{locus}_cytosine_matrix_complete.csv': matrix_csv(d)}
        new[xname] = (lcyc_workbook(old[xname], d, qc) if locus == 'Lcyc' else lvhirz_workbook(old[xname], qc))
        packages[SUPP/zname] = {k: new[k] for k in sorted(new, key=lambda k: (not k.endswith('.xlsx'), k))}
    matrix, d = tables['Lcyc']
    s7 = s7_members(read_zip(SUPP/'Data_S7_architecture_space.zip'), d, matrix)
    packages[SUPP/'Data_S7_architecture_space.zip'] = s7
    return packages

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--write', action='store_true', help='rebuild the packages instead of checking them')
    args = ap.parse_args()
    packages = build()
    lcyc_csv = next(v for k, v in packages[SUPP/'Data_S1_Lcyc.zip'].items() if k.endswith('matrix_complete.csv'))
    if args.write:
        for path, members in packages.items():
            write_zip(path, members)
        BENCH.write_bytes(lcyc_csv)
        print('Rebuilt Supplementary Data S1, S2 and S7 and the benchmark input')
        return
    ok = True
    for path, members in packages.items():
        dep = read_zip(path)
        for name, data in members.items():
            good = name in dep and same(dep[name], data, name)
            ok &= good; print('PASS' if good else 'FAIL', path.name, name.split('/')[-1])
        extra = set(dep) - set(members)
        if extra:
            ok = False; print('FAIL', path.name, 'unexpected members', sorted(extra))
    good = BENCH.read_bytes() == lcyc_csv; ok &= good; print('PASS' if good else 'FAIL', 'benchmark Peloria input')
    if not ok:
        sys.exit(1)
    print('DATA PACKAGES MATCH THE RECONSTRUCTED MATRICES')

if __name__ == '__main__':
    main()
